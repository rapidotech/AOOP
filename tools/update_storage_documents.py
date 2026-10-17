"""Append observed local storage results without rewriting historical log entries."""
from pathlib import Path
import json
import shutil
import zipfile
from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.shared import Inches, Pt, RGBColor

OUTPUT = Path.home() / 'Downloads/AOOP_Milestone2_Group13'
TITLE = 'Local Docker object storage update'


def main():
    proof = json.loads((OUTPUT/'evidence/local-object-storage-assurance.json').read_text())
    schedule_bytes = (OUTPUT/'evidence/local-object-storage-scheduler.json').read_bytes()
    schedule = json.loads(schedule_bytes.decode('utf-16' if schedule_bytes.startswith(b'\xff\xfe') else 'utf-8-sig'))
    if proof['assertions'] != 18 or proof['independent_offhost_custody'] is not False:
        raise ValueError('Unexpected assurance scope; review the evidence before editing')
    if schedule['result'] != 0 or not schedule['new_s3_receipt_created']:
        raise ValueError('Successful scheduled upload not recorded')
    target = OUTPUT/'MotorPH_Group13_Milestone2_Implementation_Log.docx'
    baseline = Path.home()/'.motorph/milestone2/object-storage/docs-baseline'
    baseline.mkdir(exist_ok=True)
    doc = Document(target)
    if any(p.text == TITLE for p in doc.paragraphs):
        raise ValueError('Addendum already exists; preserve subsequent edits')
    for extension in ('docx', 'pdf'):
        source = target.with_suffix('.'+extension)
        saved = baseline/source.name
        if saved.exists():
            if extension == 'docx' and saved.read_bytes() != source.read_bytes():
                raise ValueError('Baseline differs; refusing to overwrite historical edits')
        else:
            shutil.copyfile(source, saved)
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.page_width, section.page_height = Inches(8.27), Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(.65)
    section.left_margin = section.right_margin = Inches(.7)
    heading = doc.add_paragraph(TITLE, 'Title')
    for run in heading.runs:
        run.font.size, run.font.color.rgb = Pt(21), RGBColor(0, 0, 0)
    entries = [
        ('Current log and scope', '17 October 2026. Use outputs/implementation-log/MotorPH_Group13_Implementation_Log.xlsx as the current school-template log. Earlier pages retain implementation chronology. Their 331-test and CI-not-yet-run statements describe the earlier stage; GitHub run38997465784 subsequently passed 352 offline assertions, source and dependency gates. C1 and C2 are the original selections; this addendum records later supporting C5 storage work.'),
        ('Implemented locally', 'RustFS 1.0.2 runs in the dedicated motorph-m2-archive Docker Desktop project. S3 port 19000 and console port 19001 bind 127.0.0.1 only. The private motorph-audit-archives bucket has versioning and seven-day COMPLIANCE default retention, a lab setting rather than an approved institutional policy. Separate protected root and scoped writer credentials are retained privately.'),
        ('Observed assurance', 'Eighteen real checks passed: exact-version archive download, checksum and full-chain verification; seven AccessDenied/403 cases; corrupt-chain rejection before upload; container restart persistence; visible failure and preservation of a new local archive while S3 was unavailable. Denials included anonymous read, writer privilege/prefix misuse, and root API deletion or shortening of protected retention. The snapshot contained '+str(proof['retained_events'])+' audit events.'),
        ('Scheduled upload', 'An actual invocation of MotorPH-M2-Audit-Archive returned 0 and produced a new verified S3 receipt. The daily 23:00 task requires the provisioning Windows user signed in, Docker Desktop running and the owned MySQL clone available. Continuous operation while signed out has not been verified.'),
        ('Status and limitations', 'C5 remains Needs adjustment / Retest for its broader independent-custody requirement. The user selected storage on this same PC. Host or Docker administrators can alter the underlying volume; loss of the PC can affect both database and backup. An authorized independent host and a recovery test there are still needed. No such deployment or test is claimed.'),
        ('Evidence and access', 'See evidence/local-object-storage-assurance.json, local-object-storage-scheduler.json and local-object-storage-dependencies.json. The dependency receipt records zero OSV advisories for seven pinned Python packages at query time; no full container vulnerability scan is claimed. Setup and recovery: security/object-storage/README.md in the repository. Console: http://127.0.0.1:19001. Shared workbook access and mentor validation remain group actions, not completed test results.')]
    for title, body in entries:
        p = doc.add_paragraph(title, 'Heading 2')
        p.paragraph_format.space_before = Pt(8)
        p.paragraph_format.space_after = Pt(3)
        p = doc.add_paragraph(body)
        p.paragraph_format.space_after = Pt(5)
        for run in p.runs:
            run.font.size = Pt(10)
    temporary = target.with_suffix('.storage.docx')
    doc.save(temporary)
    # Replace only the expanded diagram image; earlier diagrams remain historical.
    with zipfile.ZipFile(temporary) as archive:
        files = {name: archive.read(name) for name in archive.namelist()}
    xml = etree.fromstring(files['word/document.xml'])
    ns = {'pic':'http://schemas.openxmlformats.org/drawingml/2006/picture',
          'a':'http://schemas.openxmlformats.org/drawingml/2006/main',
          'r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    pictures = xml.xpath('//pic:pic[pic:nvPicPr/pic:cNvPr[@name="MotorPH_Expanded_Control_Diagram.png"]]',namespaces=ns)
    if len(pictures) != 1:
        raise ValueError('Expected exactly one current diagram')
    rid = pictures[0].find('.//a:blip', ns).get('{'+ns['r']+'}embed')
    rels = etree.fromstring(files['word/_rels/document.xml.rels'])
    relationship = next(r for r in rels if r.get('Id') == rid)
    files['word/'+relationship.get('Target')] = (OUTPUT/'MotorPH_Expanded_Control_Diagram.png').read_bytes()
    with zipfile.ZipFile(target, 'w', zipfile.ZIP_DEFLATED) as archive:
        for name, data in files.items():
            archive.writestr(name, data)
    temporary.unlink()
    print('Historical log preserved; local storage addendum and current diagram saved.')


if __name__ == '__main__':
    main()
