"""Build the editable Group 13 implementation log and standalone diagram PDF.
Run with the bundled document-runtime Python, not the application environment.
"""
from pathlib import Path
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A3, landscape
from reportlab.lib.utils import ImageReader

OUTPUT = Path.home() / 'Downloads/AOOP_Milestone2_Group13'
EVIDENCE = OUTPUT / 'evidence'
DATE = '17 October 2026'
BASELINE = '8387a21043f422e186a2a03a932eb8f50ee364c2'


def read(name):
    return json.loads((EVIDENCE / name).read_text(encoding='utf-8'))


def paragraph(document, text, bold_lead=None):
    p = document.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def field(document, label, value):
    paragraph(document, label + ': ' + value, label + ':')


def table(document, headers, rows, widths):
    t = document.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for col, width in zip(t.columns, widths):
        col.width = Inches(width)
    table_width = t._tbl.tblPr.find(qn('w:tblW'))
    table_width.set(qn('w:type'), 'dxa')
    table_width.set(qn('w:w'), str(round(sum(widths) * 1440)))
    for c, text in zip(t.rows[0].cells, headers):
        c.text = text
        shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'DCE6F1'); c._tc.get_or_add_tcPr().append(shading)
        for r in c.paragraphs[0].runs: r.bold = True
    repeat = OxmlElement('w:tblHeader'); t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for index, row in enumerate(rows):
        cells = t.add_row().cells
        for c, text in zip(cells, row):
            c.text = str(text)
            if index % 2:
                shading = OxmlElement('w:shd'); shading.set(qn('w:fill'), 'F3F6F9'); c._tc.get_or_add_tcPr().append(shading)
    for row in t.rows:
        no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        for c, width in zip(row.cells, widths):
            c.width = Inches(width)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = c._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ['top', 'left', 'bottom', 'right']:
                m = OxmlElement('w:' + side); m.set(qn('w:w'), '90'); m.set(qn('w:type'), 'dxa'); margins.append(m)
            props.append(margins)
            borders = OxmlElement('w:tcBorders')
            for side in ['top', 'left', 'bottom', 'right']:
                edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9'); borders.append(edge)
            props.append(borders)
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs: run.font.size = Pt(9.5)
    return t


def page(document, title):
    document.add_page_break()
    document.add_heading(title, level=1)


def configure(section, *, wide=False):
    section.top_margin = Inches(.65); section.bottom_margin = Inches(.65)
    section.left_margin = Inches(.7); section.right_margin = Inches(.7)
    if wide:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(16.54); section.page_height = Inches(11.69)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.27); section.page_height = Inches(11.69)


def main():
    integration = read('C1-C2-isolated-integration.json')
    enrollment = read('C2-controlled-enrollment.json')
    guards = read('review-correction-regressions.json')
    baseline = read('E02-original-baseline-revalidation.json')
    hash_ms = next(x.split('=')[-1] for x in integration['outcomes'] if 'hash duration_ms=' in x)
    document = Document()
    for border in list(document.styles.element.iter(qn('w:pBdr'))):
        border.getparent().remove(border)
    configure(document.sections[0])
    for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'Heading 3']:
        style = document.styles[name]
        style.font.name = 'Calibri'; style.font.color.rgb = RGBColor(0, 0, 0)
        if style.type == 1:
            style.paragraph_format.space_after = Pt(7)
    document.styles['Normal'].font.size = Pt(11)
    document.styles['Normal'].paragraph_format.line_spacing = 1.12
    document.styles['Title'].font.size = Pt(22)
    document.styles['Heading 1'].font.size = Pt(16)
    document.styles['Heading 2'].font.size = Pt(12)
    header = document.sections[0].header.paragraphs[0]
    header.text = 'GROUP 13   |   MO IT153 INFORMATION ASSURANCE AND SECURITY 2'
    header.runs[0].font.size = Pt(8)
    footer = document.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('MotorPH Implementation Log   |   ')
    number = OxmlElement('w:fldSimple'); number.set(qn('w:instr'), 'PAGE'); footer._p.append(number)
    document.core_properties.title = 'MotorPH Milestone 2 Implementation Log'
    document.core_properties.author = 'Group 13'
    document.add_paragraph('MotorPH Milestone 2 Implementation Log', 'Title')
    paragraph(document, 'Implemented and tested in an isolated local environment on ' + DATE + '.')
    paragraph(document, 'This log connects the two controls selected in our Security Implementation Plan to the code, configuration, database changes, and assurance evidence. Both controls were verified in the isolated MotorPH clone. The original aoop database and shared root account were left unchanged.')
    field(document, 'Project and group', 'MotorPH Payroll System - Group 13')
    field(document, 'Group members', 'Tristan Marc A. Cunanan; Charles Mikael Medel; Juan Carlos Miguel Basa; Frances Balgos')
    field(document, 'Source documents', 'Group 13 Security Implementation Plan PDF, pages 4 to 7, and the supplied audit workbook Audit Worksheet, rows 27 to 38 and 73 to 74. Originals were read only and retained.')
    paragraph(document, 'No school Implementation Log template was supplied. This editable original log follows the fields required in the five Midgame Move instructions; it does not claim to be the school template.')
    document.add_heading('Control mapping', level=1)
    table(document, ['Control and goal', 'Component', 'Surface and insertion'], [
        ['C1 Database credential protection and least privilege\nReduce secret exposure and database blast radius.', 'DatabaseConfig, protected launch profile, MySQL application account.', 'Code and configuration: DPAPI launcher and TLS checks.\nDatabase control: table grants and restricted column updates.'],
        ['C2 Password storage and default password remediation\nReplace weak shared credentials and prevent normal access before reset.', 'PasswordService, AuthenticationService, login and reset UI, users table.', 'Code: BCrypt12 creation, verification, policy and reset gate.\nDatabase control: widened hash field, reset flag and transactional replacement.']
    ], [2.2, 1.85, 2.82])
    field(document, 'Scope decision', 'Exactly C1 and C2 were implemented. Lockout, per-user DAL authorization, and a new persistent audit trail remain separate controls, not silently included or claimed complete.')

    page(document, 'Safe environment and recovery readiness')
    table(document, ['Checklist item', 'Confirmed evidence'], [
        ['Full backup completed', '701 working-tree files in a protected archive; Git bundle plus working/index patches; local configuration; complete logical SQL backup with routines, triggers and events; file hashes.'],
        ['Controlled environment', 'Separate MySQL 8.0.42 process and data directory on 127.0.0.1 port 3307. Restored schema aoop_m2; original MySQL80 service and aoop on 3306 untouched.'],
        ['Change tracking initialized', 'Pre-implementation baseline committed as 8387a21. No remote push or history rewrite.'],
        ['Rollback path identified and tested', 'Baseline Git archive extracted and rebuilt. SQL restored into aoop_m2_restore_check; all 14 table row counts matched.'],
        ['Restart and isolation checked', 'Owned isolated process stopped gracefully and restarted through the DPAPI launcher; original MySQL80 service remained running. Actual profile/account probe passed.']
    ], [1.9, 4.97])
    field(document, 'Protected backup location', 'Downloads/AOOP_Milestone2_Group13/private-backup/preimplementation-20261017. It contains sensitive recovery material and is excluded from the shareable portfolio.')
    field(document, 'Baseline commit', BASELINE)
    paragraph(document, 'For recovery, preserve current code and escrow first. Extract baseline-8387a21.zip into a separate directory and rebuild it. Restore aoop-baseline.sql into a separate recovery schema, compare the recorded counts, and only then decide whether to replace the owned test environment. Never restore over the original database or automatically reactivate an exposed secret.')
    field(document, 'Evidence', 'readiness-and-restore.json; rollback-and-source-guard.json; protected backup manifest.json.')
    paragraph(document, 'The MySQL dump used a transaction-consistent logical backup and explicit routines, triggers and events options. Restore testing, rather than the presence of a .sql file alone, supports the recovery claim. [S4]')

    page(document, 'Control 1 implementation record')
    field(document, 'Control name and type', 'Database credential protection and least privilege - Technical')
    field(document, 'Security purpose', 'Keep reusable application credentials out of current source and limit what a stolen application database credential can do.')
    field(document, 'Component and surface', 'DatabaseConfig and launch scripts: code/configuration. MySQL account and grants: database/control. Loopback binding and verified TLS support the same connection boundary.')
    field(document, 'Tools and versions', 'Java 17; MySQL 8.0.42; Connector/J 8.0.33; Windows PowerShell 5.1 DPAPI CurrentUser; a dedicated CA and PKCS12 public-certificate truststore.')
    field(document, 'Files and locations', 'DatabaseConfig.java; scripts/start-milestone2.ps1; scripts/stop-milestone2.ps1; tools/milestone2_environment.py; tools/milestone2_secure_profile.py; scripts/check-security.py. Protected runtime: %USERPROFILE%/.motorph/milestone2.')
    field(document, 'Date implemented', DATE)
    document.add_heading('Initial behavior confirmation', level=2)
    paragraph(document, 'The supplied audit recorded embedded root credentials and optional secure transport. Intake source review confirmed the embedded credential. The earlier cleanup externalized it before the committed M2 baseline, but that baseline still used a plaintext user-local file and the shared root account. Those facts were not presented as a finished security control.')
    document.add_heading('Changes applied', level=2)
    paragraph(document, 'A separate database instance was restored from backup. A localhost-only application account received only the table operations used by the app, with password/reset column updates on users. A generated old isolated secret was replaced; its rejection was tested. The active credential was exported with DPAPI encryption and a user/SYSTEM-only file ACL. Plaintext values are not placed in Git, logs, or command arguments. [S2]')
    paragraph(document, 'The clone requires secure transport. The app verifies the server certificate chain and host identity using its own CA. The launcher validates the actual JDBC target, rejects inherited JVM overrides, and checks the active database/account before opening MotorPH. [S3]')
    document.add_heading('Issues encountered and resolutions', level=2)
    paragraph(document, 'The first real CRUD test exposed a missing DELETE privilege needed by employee cleanup in attendance. That single table privilege was added; payroll remains read only. Source scanners had missed the earlier literal, so a focused guard was added and proved against a synthetic embedded-secret fixture. Independent review also required explicit launcher isolation and ownership-limited test cleanup; regression checks now cover both.')
    field(document, 'Status and next step', 'Verified in aoop_m2 - Complete for this controlled implementation. Original shared-root rotation and Git-history cleanup remain a separate coordinated rollout task.')

    page(document, 'Control 1 assurance tests')
    paragraph(document, 'The three scenario groups below combine valid and invalid triggers while retaining exact console/JSON evidence. An expected block is recorded as denied, not as a failed test.')
    table(document, ['Scenario and trigger', 'Expected and actual observation', 'Status and evidence'], [
        ['C1 T1 Protected configuration\nStart with DPAPI profile; test missing/corrupt settings, JVM override, and a mislabeled URL.', 'Expected: normal launch uses only the restricted clone; unsafe inputs stop before connection.\nObserved: current-user decryption and actual account/DB probe passed. Corrupt ciphertext and missing settings were denied. JVM override and URL mismatch were rejected. Focused source guard passed and its synthetic bad fixture was detected.', 'Verified\nNext Complete\nC1-protected-config.json\nreview-correction-regressions.json\nrollback-and-source-guard.json'],
        ['C1 T2 Required access and restrictions\nRead existing views; create/update/delete a disposable employee and salary; attempt privileged operations and plaintext transport.', 'Expected: required app operations succeed; DDL/admin/cross-schema/role-column changes fail.\nObserved: employee/salary/attendance/payroll reads and synthetic employee CRUD/salary update succeeded. System-account data and cross-schema reads denied 1142; DDL denied 1142; GRANT denied 1044; role-column update denied 1143. TLS-disabled connection denied 1045.', 'Verified\nNext Complete\nC1-C2-isolated-integration.json'],
        ['C1 T3 Rotation and recovery\nReplace isolated application secret; retry old secret; restore the backup and restart the owned instance.', 'Expected: new secret works, old secret fails, data recovers without touching the source instance.\nObserved: replacement succeeded; superseded isolated secret was denied. Fourteen tables restored with matching counts. Baseline code rebuilt; isolated restart and protected profile probe succeeded; original service and counts unchanged.', 'Verified\nNext Complete\nC1-profile-and-rotation.json\nreadiness-and-restore.json\nrollback-and-source-guard.json']
    ], [1.75, 3.3, 1.82])
    paragraph(document, 'The wrong-Windows-user DPAPI test was not performed. Microsoft documents its user/machine binding; the local proof here is current-user success, unreadable corruption, absence of plaintext in the credential file, and the restricted ACL. Same-user process access remains a limit. [S2]')

    page(document, 'Control 2 implementation record')
    field(document, 'Control name and type', 'Password storage and default password remediation - Technical')
    field(document, 'Security purpose', 'Replace unsalted SHA-256 and shared defaults with salted, deliberately costly password verification; prevent normal access until a temporary credential has been changed.')
    field(document, 'Component and surface', 'PasswordService and AuthenticationService, login and mandatory reset dialog: code. users.password and users.must_change_password: database/control.')
    field(document, 'Tools and versions', 'Favre BCrypt 0.10.2, bytes 1.5.0, cost 12; Java 17 and MySQL 8.0.42. JAR hashes and license are pinned under security/dependencies. [S1]')
    field(document, 'Files and locations', 'PasswordService.java; AuthenticationService.java; DatabaseHandler.java; frm_Login.java; PasswordResetDialog.java; DataMigration.java; header-only CSV_Users.csv; security/migrations/001_expand_password_reset.sql; security/tools/SecurityAssurance.java; scripts/protect-bootstrap-escrow.ps1.')
    field(document, 'Date implemented', DATE)
    document.add_heading('Initial behavior confirmation', level=2)
    paragraph(document, f'The unchanged original database was rechecked using aggregate queries only: {baseline["user_count"]} users, {baseline["unique_password_records"]} unique stored password value, {baseline["legacy_sha256_count"]} SHA-256-format records, and {baseline["documented_default_match_count"]} controlled matches to the documented default. No raw password/hash values were captured. This revalidation supports E02 and does not claim the original database was migrated.')
    document.add_heading('Changes applied', level=2)
    paragraph(document, 'The isolated schema was expanded idempotently. All 34 legacy accounts were given independently generated temporary plaintext credentials, hashed directly to BCrypt12 and marked for mandatory change. No SHA digest was wrapped into BCrypt. New account creation also writes BCrypt12 immediately. The known shared default cannot log in or claim a reset.')
    paragraph(document, 'Authentication returns no employee or role for a flagged account. Reset revalidates the temporary password and reset state under a row lock, checks password policy, writes the replacement and clears the flag in one transaction, and returns full authentication only after commit. Cancelling or failing reset leaves no normal session. The policy is at least 12 Unicode characters and at most 72 UTF-8 bytes, with explicit rejection instead of truncation. [S1]')
    document.add_heading('Issues encountered and resolutions', level=2)
    paragraph(document, 'A shared known default was rejected as enrollment proof; unique administrator-delivered temporary credentials were chosen instead. Review required atomic escrow replacement before committing enrollment and preservation of old ciphertext. Windows PowerShell 5 newline/array serialization problems were found during real credential retrieval, recovered without losing the 34 pairs, and fixed by preserving canonical parent JSON. Retrieval and actual account proofs were then rechecked.')
    field(document, 'Status and next step', 'Verified in aoop_m2 - Complete. Each authorized user completes the required password change using their separately delivered temporary credential.')

    page(document, 'Control 2 assurance tests')
    table(document, ['Scenario and trigger', 'Expected and actual observation', 'Status and evidence'], [
        ['C2 T1 Creation and policy\nCreate a disposable user; attempt short, known-default/repeated, and oversized replacements.', 'Expected: new storage uses unique salted BCrypt12; unacceptable passwords are rejected without truncation.\nObserved: cost-12 $2b$ records were stored; repeated hashing of the same input produced different salts. Short/default/73-byte inputs were denied. Offline tests also cover Unicode and malformed input.', 'Verified\nNext Complete\nC1-C2-isolated-integration.json\noffline-verification.txt'],
        ['C2 T2 Correct and incorrect credentials\nUse a correct temporary password, a wrong password, and the old temporary value after reset.', 'Expected: incorrect credentials fail, temporary proof grants reset only, replacement works after commit.\nObserved: wrong/default passwords denied; pending result had no employee/role. Reset committed valid BCrypt and flag 0; new password worked and old temporary password failed.', 'Verified\nNext Complete\nC1-C2-isolated-integration.json'],
        ['C2 T3 Enrollment and mandatory reset\nMigrate legacy clone accounts; cancel/fail/reset; retry migration and protected escrow.', f'Expected: no normal session before successful reset; safe retries preserve accounts and recovery secrets.\nObserved: 34 accounts enrolled; already supported BCrypt rows skipped on retry. Cancel/error UI stayed restricted. Real reset revalidated current proof and failed attempts left the flag set. Atomic/canonical escrow checks and protected retrieval passed. One measured cost-12 hash took {hash_ms} ms on this host.', 'Verified\nNext Complete\nC2-controlled-enrollment.json\nC2-escrow-atomic-maintenance.json\nreview-correction-regressions.json\nui/ui-proof.json']
    ], [1.7, 3.35, 1.82])
    paragraph(document, 'Timing is a single observed hash operation, not a performance guarantee. The selected cost matched the audit proposal and was feasible in the asynchronous local UI. Real database proofs and injected UI proofs are recorded separately.')

    page(document, 'Password reset visual evidence')
    paragraph(document, 'The mandatory reset screen keeps temporary-password proof separate from the new password. It provides confirmation, visible policy, asynchronous saving, and a route back to sign in before submission. No normal session is available while reset is required.')
    picture = document.add_paragraph()
    picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture.add_run().add_picture(str(EVIDENCE / 'ui/reset-default.png'), width=Inches(3.65))
    paragraph(document, 'Figure 1  Reset interface rendered from an injected fixture. This image contains no real credential and proves presentation only. Real storage, flag transitions, and authentication outcomes are supported by the isolated integration tests.')
    field(document, 'Additional captured states', 'Password length rejection, UTF-8 byte-limit rejection, save pending, and incorrect current password. Five sanitized images and hashes are in evidence/ui. Two bounded inspection rounds were used.')
    field(document, 'Integrated offline checks', '75 persistence, 53 security, and 63 UI assertions passed. Runtime tests exclude the real MySQL driver. The isolated database suite is separate: ' + str(integration['assertions']) + ' checks; reviewer correction regressions: ' + str(guards['assertions']) + '.')

    wide = document.add_section(WD_SECTION_START.NEW_PAGE); configure(wide, wide=True)
    document.add_heading('Updated system diagram', level=1)
    paragraph(document, 'Figure 2  Current control integration in the isolated MotorPH environment. Added and modified labels identify C1 and C2 surfaces; the original database is outside the changed boundary.')
    document.add_picture(str(OUTPUT / 'MotorPH_Updated_Control_Diagram.png'), width=Inches(15.1))
    paragraph(document, 'C1 adds protected configuration, restricted grants, and verified TLS. C2 modifies authentication and adds the mandatory reset gate. The reset service uses the same transactional JDBC path; a session is released only after commit. The normal employee/payroll workflows remain behind successful authentication. Local attendance request/payroll ledgers retain their separate documented storage boundary.')

    normal = document.add_section(WD_SECTION_START.NEW_PAGE); configure(normal)
    document.add_heading('Assessment and evidence boundaries', level=1)
    paragraph(document, 'The selected controls behaved as intended in the isolated environment. Expected blocks were observed, not inferred from source alone. Recovery was rehearsed, and independent source review corrections were implemented and regression-tested. This is not evidence of production remediation or of every original audit risk being closed.')
    table(document, ['Area', 'Observed status or remaining boundary'], [
        ['C1 and C2', 'Verified in the dedicated clone, including state-changing tests on disposable records.'],
        ['Original root and Git history', 'Unchanged. The exposed original shared credential needs coordinated rotation and history remediation before real rollout.'],
        ['Other audit controls', 'Lockout, persistent audit trail, and per-user DAL authorization remain deferred. Denying role-column UPDATE does not prove complete user authorization.'],
        ['Configuration compatibility', 'The plaintext legacy fallback still exists for compatibility. The verified active launch profile uses DPAPI, not that fallback. Same-Windows-user compromise remains outside DPAPI protection.'],
        ['Other system proof', 'Physical printing, statutory-rate correctness, internet exposure, production security and multi-process local-ledger locking were not verified. Existing print/filter/navigation paths were preserved.']
    ], [1.7, 5.17])
    document.add_heading('Code and architecture cleanup', level=2)
    paragraph(document, 'Before the security work, startup moved to the Swing event thread; build/library paths were standardized for Java 17; configuration, transactions, storage, authentication and shared UI responsibilities were separated. Atomic local records and exact-record approvals fixed misleading save behavior. All 16 forms gained shared accessibility/readability treatment. Three hundred ninety generated output files were removed from version control while local copies were preserved.')
    document.add_heading('Ongoing implementation log use', level=2)
    paragraph(document, 'Keep this DOCX and the repository security notes as the working log. Add new dated entries and evidence rather than replacing these observations. Further deployments, privilege changes, or shared-root remediation need their own environment-specific tests and status. No mentor approval or unperformed group-member activity is asserted.')

    page(document, 'Evidence index and reproduction')
    table(document, ['Evidence file', 'What it supports'], [
        ['readiness-and-restore.json', '14-table restore and original counts unchanged'],
        ['rollback-and-source-guard.json', 'Baseline rebuild, restart, positive and negative scanner checks'],
        ['C1-profile-and-rotation.json', 'Actual grants and isolated secret rotation'],
        ['C1-protected-config.json', 'DPAPI/current-user/corrupt-input observations and limits'],
        ['C1-C2-isolated-integration.json', str(integration['assertions']) + ' real isolated database checks and SQL denial codes'],
        ['C2-controlled-enrollment.json', '34 initial enrollments and retry-safe BCrypt skipping'],
        ['C2-escrow-atomic-maintenance.json', 'Protected map recovery, canonical storage and retrieval'],
        ['review-correction-regressions.json', str(guards['assertions']) + ' focused safety regressions'],
        ['ui/ui-proof.json and PNGs', 'Injected UI states and rendering provenance'],
        ['offline-verification.txt', 'Final executable regression suite output']
    ], [3.25, 3.62])
    paragraph(document, 'Run scripts/build.ps1 -Test, scripts/check-security.py, and scripts/start-milestone2.ps1 -SkipBuild -VerifyOnly. The controlled integration commands are tools/test_milestone2_guards.py and tools/run_milestone2_assurance.py. Run only on the marked prepared clone; security/README.md contains prerequisites and recovery instructions.')
    document.add_heading('Source references', level=2)
    for text in [
        'SIP and audit  Supplied Group 13 PDF pages 4 to 7 and audit workbook rows 27 to 38 and 73 to 74. Original file hashes: evidence/source-document-hashes.json.',
        'S1 BCrypt API and limits  https://github.com/patrickfav/bcrypt',
        'S2 Windows DPAPI credential export  https://learn.microsoft.com/en-us/powershell/module/Microsoft.PowerShell.Utility/export-clixml?view=powershell-5.1',
        'S3 MySQL encrypted connections  https://dev.mysql.com/doc/refman/8.0/en/using-encrypted-connections.html',
        'S4 MySQL backup behavior  https://dev.mysql.com/doc/refman/8.0/en/mysqldump.html'
    ]:
        p = paragraph(document, text)
        for r in p.runs: r.font.size = Pt(9)
    path = OUTPUT / 'MotorPH_Group13_Milestone2_Implementation_Log.docx'
    document.save(path)
    pdf = OUTPUT / 'MotorPH_Updated_Control_Diagram.pdf'
    width, height = landscape(A3)
    c = canvas.Canvas(str(pdf), pagesize=(width, height))
    c.setTitle('MotorPH Milestone 2 Control Integration')
    c.setFont('Helvetica-Bold', 21); c.drawString(40, height-48, 'MotorPH Milestone 2 Control Integration')
    c.setFont('Helvetica', 11); c.drawString(40,height-70,'Group 13  |  Implemented and tested in the isolated local environment  |  17 October 2026')
    image_width=width-80; image_height=image_width*660/1390
    c.drawImage(ImageReader(str(OUTPUT/'MotorPH_Updated_Control_Diagram.png')),40,height-100-image_height,width=image_width,height=image_height)
    c.setFont('Helvetica', 11)
    y=height-125-image_height
    for line in [
        'C1  Added DPAPI configuration, least-privilege application grants, credential rotation and verified TLS.',
        'C2  Added BCrypt cost 12 and a mandatory reset gate; no normal session before reset commits.',
        'The reset service uses the same transactional JDBC path. The original aoop database and shared root are unchanged.',
        'Evidence and limitations are recorded in the Group 13 Implementation Log. This is isolated proof, not production assurance.'
    ]:
        c.drawString(40,y,line); y-=18
    c.showPage();c.save()
    print('Created editable Implementation Log and annotated diagram PDF.')


if __name__ == '__main__':
    main()
