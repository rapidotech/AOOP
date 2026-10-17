"""Append observed R2 results to the chronological log; preserve prior entries."""
from pathlib import Path
import json,zipfile
from lxml import etree
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.shared import Inches,Pt,RGBColor

OUTPUT=Path.home()/'Downloads/AOOP_Milestone2_Group13'
TITLE='Cloudflare R2 off host backup update'


def main():
    proof=json.loads((OUTPUT/'evidence/cloudflare-r2-assurance.json').read_text())
    schedule=json.loads((OUTPUT/'evidence/cloudflare-r2-scheduler.json').read_text(encoding='utf-8-sig'))
    if proof['assertions']!=20 or not proof['offhost_storage_verified'] or schedule['result']!=0:
        raise ValueError('Observed live R2 and scheduled evidence required')
    path=OUTPUT/'MotorPH_Group13_Milestone2_Implementation_Log.docx'
    doc=Document(path)
    if any(p.text==TITLE for p in doc.paragraphs):raise ValueError('R2 addendum already present; preserve later edits')
    section=doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.page_width,section.page_height=Inches(8.27),Inches(11.69)
    section.top_margin=section.bottom_margin=Inches(.65)
    section.left_margin=section.right_margin=Inches(.7)
    heading=doc.add_paragraph(TITLE,'Title')
    for run in heading.runs:run.font.size,run.font.color.rgb=Pt(21),RGBColor(0,0,0)
    entries=[
      ('Current assessment','17 October 2026. This later addendum supersedes earlier statements that off-host copies were not provisioned. The current school-template workbook records C5 as Verified / Complete for its tested transactional audit and recoverable local/off-host copies. C1/C2 remain the two original SIP selections. Independent administration is outside the verified claim.'),
      ('Implemented surfaces','The private motorph-audit-archives R2 bucket uses Standard storage and HTTPS. Public r2.dev access is disabled; no public custom domain exists. A bucket-scoped object writer uploads archives; a separate read-only token checks configuration. Generated tokens expire after90days and are protected with DPAPI CurrentUser. The existing local files and RustFS Docker copies remain.'),
      ('Actual assurance','Twenty live checks passed, including download of an actual302-event audit snapshot, matching checksum/full chain and protected manifest. Eight operations were rejected: anonymous read; writer delete/overwrite; bootstrap S3 deletion while locked; two configuration changes; and read/write access to a separate owned test bucket. Provider-specific errors are recorded, including missing Authorization HTTP400 and ObjectLockedByBucketPolicy.'),
      ('Schedule and retention','An actual scheduled invocation returned0 and produced verified Docker/R2 receipts with the same archive checksum. The task is daily23:00 while this Windows user is signed in. Seven-day R2 bucket locks are administrator-removable. Eight-day expiry is configured but its future deletion has not yet been observed. Neither R2 S3 COMPLIANCE retention nor independently administered custody is claimed; the bootstrap credential remains on this PC.'),
      ('Metered free tier safeguards','The user accepted metered free-tier use after the billing limitation was explained. Limits are10MiB per archive,256MiB in the owned bucket,200 request attempts/day and2000/month including SDK retries. Corrupt or exhausted counters fail closed. No paid add-on, Worker or public endpoint was enabled. These limits are intended to stay within the free allowance; no provider hard $0 cap or zero-dollar invoice is certified.'),
      ('Evidence and remaining actions','See cloudflare-r2-assurance.json, cloudflare-r2-scheduler.json, cloudflare-r2-offline-safety.json and CLOUDFLARE_R2.md. Offline safety tests use fixtures and are separate from the20 live checks. Shared log upload/access, Google Sheets conversion and mentor validation are still group actions. Credentials and actual audit objects are excluded from the shared package.')]
    for title,body in entries:
        p=doc.add_paragraph(title,'Heading 2');p.paragraph_format.space_before=Pt(7);p.paragraph_format.space_after=Pt(3)
        p=doc.add_paragraph(body);p.paragraph_format.space_after=Pt(4)
        for run in p.runs:run.font.size=Pt(10)
    temp=path.with_suffix('.r2.docx');doc.save(temp)
    with zipfile.ZipFile(temp) as archive:files={n:archive.read(n) for n in archive.namelist()}
    xml=etree.fromstring(files['word/document.xml'])
    ns={'pic':'http://schemas.openxmlformats.org/drawingml/2006/picture','a':'http://schemas.openxmlformats.org/drawingml/2006/main','r':'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
    pictures=xml.xpath('//pic:pic[pic:nvPicPr/pic:cNvPr[@name="MotorPH_Expanded_Control_Diagram.png"]]',namespaces=ns)
    if len(pictures)!=1:raise ValueError('Expected exactly one current diagram')
    rid=pictures[0].find('.//a:blip',ns).get('{'+ns['r']+'}embed')
    rels=etree.fromstring(files['word/_rels/document.xml.rels'])
    target=next(r.get('Target') for r in rels if r.get('Id')==rid)
    files['word/'+target]=(OUTPUT/'MotorPH_Expanded_Control_Diagram.png').read_bytes()
    with zipfile.ZipFile(path,'w',zipfile.ZIP_DEFLATED) as archive:
        for n,data in files.items():archive.writestr(n,data)
    temp.unlink()
    print('R2 chronology addendum and current diagram saved; earlier entries retained.')


if __name__=='__main__':main()
