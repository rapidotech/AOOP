"""Append evidence-backed C3-C5 implementation records to the original C1/C2 log."""
from pathlib import Path
import json
import re
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.section import WD_SECTION_START
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A2, landscape
from reportlab.lib.utils import ImageReader
import build_milestone2_portfolio as base

OUTPUT = base.OUTPUT
EVIDENCE = base.EVIDENCE


def main():
    base.main()
    path = OUTPUT / 'MotorPH_Group13_Milestone2_Implementation_Log.docx'
    document = Document(path)
    integration = json.loads((EVIDENCE / 'expanded-controls-integration.json').read_text())
    runtime = integration['runtime']['assertions']
    sql_checks = integration['sql_checks']
    raw = (EVIDENCE / 'expanded-offline-verification.txt').read_bytes()
    verification = raw.decode('utf-16') if raw.startswith(b'\xff\xfe') else raw.decode('utf-8-sig')
    counts = {name: int(number) for name, number in re.findall(r'^([^\r\n]+?) regression checks passed: (\d+) assertions', verification, re.MULTILINE)}
    if not {'Persistence', 'Security', 'UI', 'Audit UI'}.issubset(counts):
        raise ValueError('Complete final offline evidence is required')
    total = sum(counts.values())
    document.paragraphs[1].insert_paragraph_before('Expanded revision: C1 to C5 and supporting controls. The initial C1/C2 implementation notes are retained below; the expanded records that follow supersede their earlier scope and test status. Original Milestone 1 source documents remain unchanged.')
    for p in document.paragraphs:
        if p.text.startswith('This log connects the two controls'):
            p.text = 'Initial delivery record: the following C1/C2 notes connect the two originally selected SIP controls to their actual implementation and isolated assurance evidence. The later expanded revision adds the other worthwhile audit proposals.'
        if p.text.startswith('Scope decision: Exactly C1 and C2'):
            p.text = 'Initial scope decision: C1/C2 were implemented first. C3/C4/C5 were deferred at that delivery and are implemented in the expanded revision below.'
        if p.text.startswith('Run scripts/build.ps1 -Test, scripts/check-security.py'):
            p.text = 'These initial C1/C2 records are historical. Use the current commands in the expanded reproduction section. The old enrollment runner now delegates runtime assurance to the expanded suite when migration 002 is present.'
    for t in document.tables:
        for row in t.rows:
            for cell in row.cells:
                if cell.text.startswith('Lockout, persistent audit trail, and per-user DAL authorization remain deferred.'):
                    cell.text = 'At initial delivery these controls were deferred. The expanded revision below supersedes that status. Shared direct-JDBC and independent-custody limits remain.'
    page, field, paragraph, table = base.page, base.field, base.paragraph, base.table
    page(document, 'Expanded scope and control decisions')
    field(document, 'Date and baseline', '17 October 2026. Source and isolated database preserved before expansion at baseline 9ba48e6. The initial public portfolio is retained in the protected expansion backup.')
    paragraph(document, 'The user requested assessment of every proposed control and autonomous implementation of the worthwhile feasible additions. The SIP selected C1/C2; audit worksheet rows 75 to 83 proposed the wider set. This revision records that scope change rather than rewriting the original plan or inventing mentor approval.')
    table(document, ['Proposal', 'Decision and reason', 'Current boundary'], [
        ['C1 credential protection', 'Keep; protect reusable secrets and limit SQL privileges.', 'DPAPI, restricted clone account, TLS, tuned secret scanning.'],
        ['C2 password remediation', 'Keep; salted costly verification and mandatory reset remain necessary.', 'BCrypt12, unique temporary credentials, reset gating.'],
        ['C3 lockout and notification', 'Implement; reduce unrestricted guessing and surface failures.', 'Persistent account/source throttles, audit events, HR desktop activity.'],
        ['C4 DAL/object authorization', 'Implement; enforce permissions beyond hidden buttons.', 'Verified session capability, HR/self scope, constructor guards.'],
        ['C5 audit and workflow storage', 'Implement local mechanisms; eliminate unaudited file writes.', 'Atomic mutation/audit, serialized chain, DB ledgers and local archives.'],
        ['Transport/build/repository/release', 'Keep and extend; prevent regressions and vulnerable dependencies.', 'Loopback/TLS, Java 17, pinned libraries, scanners, commit/CI/release checks.'],
        ['Off-host/email/new roles', 'Off-host remains incomplete; email and new roles stay deferred.', 'No authorized external archive/mail infrastructure or role policy supplied.']
    ], [1.65, 2.65, 2.57])
    field(document, 'Assessment reference', 'security/EXPANDED_CONTROL_PLAN.md and the revised control audit workbook. Every effectiveness statement is limited to observed local evidence.')

    page(document, 'Control 3 implementation record')
    field(document, 'Control name and type', 'Lockout, progressive delay, authentication logging and desktop notification - Technical')
    field(document, 'Goal', 'Limit repeated password guesses, retain rejected attempts and make accumulated failures visible to HR.')
    field(document, 'Component and surface', 'AuthenticationService and reset proof: code. users and AuthThrottle: database/control. AuditService and HR security activity: service/UI.')
    field(document, 'Tools and locations', 'Java 17, JDBC/MySQL 8.0.42, BCrypt 0.10.2; AuthenticationService.java, AuthenticationThrottledException.java, AuditService.java, SecurityActivityPanel.java, frm_SecurityEvents.java, migration 002.')
    field(document, 'Initial behavior', 'The original audit recorded no lockout fields, no delays and no persistent authentication log. The initial C1/C2 delivery still lacked this control.')
    paragraph(document, 'Account and source counters now persist in the database. Failed attempts receive 1/2/4/8-second retry deadlines; the fifth failure sets a 15-minute lock. The backend checks deadlines before credential acceptance and commits counters and audit events before reporting rejection. Correct credentials cannot clear an active lock. Password-reset proof uses the same enforcement. Unknown/unsupported accounts receive cost 12 dummy verification to reduce an obvious work-factor timing difference.')
    paragraph(document, 'The Swing worker remains asynchronous. Login/reset show generic retry guidance with a timer, without sleeping on the event thread. The HR activity panel and owned event dialog fetch the latest 100 failure/lockout events, including failures accumulated before HR signed in. Timestamps are explicitly UTC. No email delivery is claimed.')
    field(document, 'Issues and resolution', 'Throwing inside the transaction would discard counters, so rejection is returned internally and raised after commit. Source-machine labels are client-controlled telemetry; account-wide lockout still survives source changes.')
    field(document, 'Date and status', '17 October 2026. Verified in the supported client and isolated clone. Next step: complete for this environment.')

    page(document, 'Control 3 assurance tests')
    table(document, ['Scenario', 'Expected and observed', 'Evidence/status'], [
        ['Valid credential; ordinary failure', 'Valid disposable HR/employee credentials establish full sessions. Incorrect proofs persist failure events and positive retry deadlines; immediate retries are denied.', 'Expanded runtime suite; verified.'],
        ['Five failed attempts and reset proof', 'Retry times 1/2/4/8/900 seconds observed. SQL stores five failures and about 900 seconds remaining. Correct password during lock is denied. Incorrect reset proof and immediate correct retry are both throttled.', 'expanded-controls-integration.json; verified.'],
        ['Persistence and administrator visibility', 'A fresh JVM using a different source label still receives the account lock. HR can read accumulated alerts; employee access is denied. UI fixtures prove pending, empty, unavailable and populated presentation.', 'Real runtime plus separately labeled UI evidence; verified.']
    ], [1.7, 3.35, 1.82])
    paragraph(document, 'To keep testing bounded, the test administrator expires the disposable account/source retry deadlines between independent scenarios. The fifteen-minute duration is checked directly in SQL; the suite does not claim to have waited fifteen minutes. A second JVM provides process-boundary evidence. The actual routine restart and external device identity are separate claims.')
    field(document, 'Post-testing status and next step', 'Verified. Complete for the isolated environment. Preserve generic UI messages and the documented shared-client/source-identity limits.')

    page(document, 'Control 4 implementation and permission matrix')
    field(document, 'Control name and type', 'Data-access-layer authorization and object-level scope - Technical')
    field(document, 'Component and surface', 'AuthorizationService, EmpUserSession, DatabaseHandler, DatabaseLedger, CSV compatibility adapters and restricted form constructors: code surface.')
    field(document, 'Tools and locations', 'Java 17/JDBC. AuthorizationService.java, EmpUserSession.java, DatabaseHandler.java, DatabaseLedger.java and restricted frm_*.java constructors; migration 002 supplies supporting database state.')
    field(document, 'Initial behavior', 'The original audit found role checks in UI menus but no confirmed DAL or constructor enforcement. Caller-supplied user/role assignment was not adequate authentication proof.')
    paragraph(document, 'A successful committed authentication result issues a single-use capability. Reset-only or fabricated results cannot establish sessions. Authorization captures immutable identity for each operation and rechecks database username, role, credential fingerprint and reset state. Sensitive writes require HR; employee collection queries bind the signed-in employee ID in SQL. Guards execute before restricted forms initialize or load data. The bulk credential-dump entry point is denied.')
    table(document, ['Operation', 'Employee', 'HR'], [
        ['Profile/salary/attendance/payroll reads', 'Own employee records only', 'Administrative collection access'],
        ['Attendance request', 'Create/edit own Pending request', 'Manage requests; exact permitted approvals'],
        ['Employee/account/salary writes', 'Denied', 'Allowed with current verified HR identity'],
        ['Payroll processing and status', 'Denied', 'Allowed with enforced workflow transitions'],
        ['Security activity and legacy import', 'Denied', 'Allowed; imported data validated/audited']
    ], [2.35, 2.25, 2.27])
    field(document, 'Issues and resolution', 'Existing enum role Employee is normalized/compared safely. A mutable displayed employee object cannot supply authorization identity. Expired account state returns the user to sign-in; ordinary employee-to-HR denial preserves the current page.')
    field(document, 'Date and status', '17 October 2026. Supported-client enforcement verified. Direct use of the shared JDBC credential remains outside this application-layer boundary.')

    page(document, 'Control 4 assurance tests')
    table(document, ['Scenario', 'Expected and observed', 'Evidence/status'], [
        ['Valid HR and employee sessions', 'Genuine credentials issue consumable capabilities. HR retains collection/CRUD access. Employee profile, salary and workflow reads return self-scoped data.', 'Real isolated runtime; verified.'],
        ['Bypass attempts', 'Unauthenticated DAL calls, caller-supplied roles, reused/reset-only capabilities, employee HR deletion/account creation/activity/import and another employee request are denied. Constructor guards reject before data reads.', 'Runtime plus offline constructor/fixture tests; verified.'],
        ['Account/session changes', 'Setting must_change_password after login invalidates the existing session on its next DAL operation. Captured transaction identity is used for audit attribution.', 'Runtime and focused review; verified.']
    ], [1.7, 3.35, 1.82])
    paragraph(document, 'These controls protect the supported desktop application. Because the desktop still connects directly with a shared database account, a modified client can issue SQL permitted to that account. A trusted server/API with independent user authentication is the architectural follow-on; no server-side human-identity guarantee is claimed here.')
    field(document, 'Post-testing status and next step', 'Verified in supported-client scope. Complete for current roles. Expanded Payroll/Admin role design remains deferred because no new organizational role policy was supplied.')

    page(document, 'Control 5 implementation record')
    field(document, 'Control name and type', 'Transactional tamper-evident audit and complete workflow-write coverage - Technical')
    field(document, 'Goal', 'Retain attributable security/business events, reject unaudited writes and detect altered chain records or archive copies.')
    field(document, 'Component and surface', 'AuditService/DatabaseLedger: code. AuditLog, private chain head, append procedure, AttendanceRequests and PayrollLedger: database/control. Archive verifier and daily local task: utility/platform.')
    field(document, 'Tools and locations', 'Java 17, MySQL 8.0.42 and SHA256. Migration 002, AuditService.java, DatabaseLedger.java, AuditEvent.java, AuditTrailDialog.java, archive_audit.py and archive task scripts.')
    field(document, 'Initial behavior', 'The audit found no persistent audit table. Attendance requests and payroll workflow records used files, outside database transaction and audit coverage.')
    paragraph(document, 'The server-owned append procedure locks a private chain head, assigns a contiguous ID and hashes the previous hash, UTC time and canonical fields. It never commits independently. The app has EXECUTE/SELECT and no direct audit/head DML. Business mutation and append share one transaction; audit failure rolls back both. Before/after status metadata is allowlisted. Credentials, hashes and full payroll/personnel values are excluded.')
    paragraph(document, 'Existing CSVHandler entry points now adapt to database workflow ledgers. Employee/date/transaction keys and statuses are normalized; validated JSON arrays preserve the original20/22 fields. Exact-row approvals enforce Pending-to-Approved/Rejected attendance and Pending-to-Verified-to-Approved payroll, with supported undo to Pending. Ordinary save/edit cannot bypass approvals or erase history. Numeric input must parse to finite, domain-valid decimal values. Local write APIs are closed; obsolete unaudited migration SQL is retired.')
    paragraph(document, 'Historical import validates rows, skips exact duplicates and rolls back conflicts while preserving source files. No external local ledgers existed here; fixtures exercised both outcomes. HR can read business metadata in a separate paginated Audit trail dialog with UTC times and Newest/Older navigation. Employees are denied. Authentication alerts remain a separate view.')
    field(document, 'Issues and resolution', 'Review corrected orphan-tail verification, global-session actor reuse, status bypasses and nonfinite numbers. The private chain head is updated only by the definer routine; complete chain/head verification uses a consistent snapshot.')
    field(document, 'Date and status', '17 October 2026. Local transactional mechanisms verified. Independent off-host custody remains incomplete.')

    page(document, 'Control 5 assurance tests and archive custody')
    table(document, ['Scenario', 'Expected and observed', 'Evidence/status'], [
        ['Valid mutation and forced audit failure', 'Request/approval/import writes create audit events. Revoking routine EXECUTE makes an employee update fail; original phone and audit head remain unchanged. Privilege is restored in finally.', 'Expanded real runtime; verified.'],
        ['Concurrency and tampering', 'Eight concurrent appends produce one valid contiguous chain. Direct audit INSERT/UPDATE/DELETE and head UPDATE are denied. Archive verifier rejects edited entries, removed tails and rewound heads.', 'SQL/chain suite and optimized-mode negatives; verified.'],
        ['Migration and workflow integrity', '20/22-field import round-trips; exact retry does not duplicate. Conflicting import rolls back. Unauthorized approval, skipped verification and nonfinite values are rejected.', 'Real runtime plus backend regressions; verified.']
    ], [1.7, 3.35, 1.82])
    paragraph(document, 'archive_audit.py exports a verified consistent snapshot with SHA256 manifest into the protected runtime directory. Prior archives are not overwritten. An owned daily 23:00 task was registered for the provisioning Windows user and successfully executed. It runs while that user is signed in; the database must be available. Optimized Python cannot disable verification or dependency checks.')
    paragraph(document, 'No authorized independent external archive destination was provided. The optional explicit destination argument can copy and checksum an export, but no off-host transfer/custody is claimed. A database administrator can rewrite both database history and local anchors. The full off-host proposal therefore remains Partial even though local logging, transactions and export mechanisms are verified.')
    field(document, 'Post-testing status and next step', 'Local mechanisms verified. Independent custody needs an authorized off-host destination and its own observed delivery/retention test; do not mark that part complete.')

    page(document, 'Supporting controls and review corrections')
    field(document, 'Network and transport', 'Clone binds 127.0.0.1:3307, requires TLS and uses VERIFY_IDENTITY. Updated driver negotiation and plaintext rejection were tested. Original port 3306 service/firewall remain unchanged.')
    field(document, 'Repository prevention', 'Custom Gitleaks rule detects low-entropy DatabaseConfig credentials. Positive detection, redacted reporting and safe external-config cases passed. A local pre-commit hook is enabled; source CI is authored. Historical secrets/public repository visibility were not changed.')
    field(document, 'Dependency correction', 'Live OSV identified GHSA-m6vm-37g8-gqvh in Connector/J 8.0.33. Version 9.7.0 was retrieved from Maven Central, checksum verified and tested with MySQL 8.0.42/Java 17. Its driver class targets Java 8. Version 26.7 requires a newer server and was not chosen. Four active JARs have a pinned SHA256 inventory; current live OSV queries returned no advisories.')
    field(document, 'Tool compatibility', 'WSL Ubuntu failed with E_UNEXPECTED, so Semgrep/Dependency-Check were not successfully rerun here. Working Gitleaks, manual review, meaningful runtime tests and live OSV were used. No NVD-completeness or vulnerability-free claim is made.')
    field(document, 'Other proposed options', 'Plaintext-verified transparent rehash is legitimate but superseded by completed unique-bootstrap resets for these known default accounts; SHA-digest wrapping is rejected. Password-column encryption with a desktop-held key and HikariCP are deferred because they add key/pool management without a demonstrated benefit in this build. Audit old/new values are limited to workflow statuses; full sensitive record dumps are excluded.')
    field(document, 'Build and release', 'Java 17 target and UTF-8 remain enforced. scripts/verify.ps1 runs offline regressions, source/secret and dependency checks; -Integration adds the marked-clone runtime suite. Authored CI covers source/build checks and has not been remotely executed.')
    field(document, 'Independent review', 'Early review corrected full-chain snapshot/head coverage and captured audit identity. Final review corrected DAL state-transition bypass, invalid numeric persistence and Python assert-based integrity guards. Optimized-mode malformed archive and changed dependency both fail after correction. Final focused review is recorded separately.')
    field(document, 'Current check totals', f'{total} offline assertions ({counts["Persistence"]} persistence, {counts["Security"]} security, {counts["UI"]} existing UI, {counts["Audit UI"]} audit UI); {runtime} real application checks and {sql_checks} SQL/chain checks. UI rendering proof is separate from real database effects.')

    section=document.add_section(WD_SECTION_START.NEW_PAGE);base.configure(section,wide=True)
    document.add_heading('Updated expanded system diagram',level=1)
    paragraph(document,'The current diagram adds C3 lockout/HR activity, C4 identity and object scope, and C5 transactional workflow/audit storage. The local archive is explicitly labeled with its remaining off-host boundary.')
    p=document.add_paragraph();p.alignment=1
    p.add_run().add_picture(str(OUTPUT/'MotorPH_Expanded_Control_Diagram.png'),width=Inches(14.6))
    paragraph(document,'The HTML provides a searchable, zoomable view. PNG/SVG and the standalone PDF retain control labels. The original C1/C2 diagram earlier in this log is historical; this expanded diagram is authoritative for the current implementation.')
    section=document.add_section(WD_SECTION_START.NEW_PAGE);base.configure(section)
    document.add_heading('Final assessment and current reproduction',level=1)
    paragraph(document,'All five core controls are worthwhile. C1/C2 remain in place; C3/C4 and the local C5 mechanisms now have runtime evidence. Supporting transport, repository, dependency and build/release checks reduce regression risk. The original source documents, database, accounts and shared service remain unchanged.')
    table(document,['Remaining boundary','Observed status'],[
        ['Original root/history/privacy','Coordinated rollout remains required. The clone uses a separate rotated account.'],
        ['Shared direct JDBC','A modified client is outside the supported-client authorization guarantee.'],
        ['Independent archive custody','Local verification/task proven. Actual off-host destination and retention are not provisioned.'],
        ['Scanner/CI completeness','Live OSV and local Gitleaks ran. New Semgrep/NVD run and remote CI are not claimed.'],
        ['Other product scope','Current roles preserved. Email/new roles, physical printing, internet exposure and statutory-rate verification remain outside these proofs.']
    ],[2.1,4.77])
    paragraph(document,'Reproduce on the prepared machine: scripts/build.ps1 -Test; python scripts/check-security.py; python scripts/check-repository.py; python scripts/check-dependencies.py; python tools/run_expanded_assurance.py; python tools/test_integrity_gates.py; python tools/test_repository_controls.py; python tools/test_milestone2_guards.py; python tools/archive_audit.py. scripts/verify.ps1 -Integration is the combined release command.')
    paragraph(document,'Use run.bat for normal protected launch. Retrieve a temporary credential with scripts/get-milestone2-bootstrap.ps1 -Username <authorized account> -CopyToClipboard, then complete the mandatory reset. security/README.md explains provisioning, import, rollback and task ownership.')

    page(document,'Expanded evidence index and revised audit worksheet')
    table(document,['Evidence','Meaning'],[
        ['expanded-offline-verification.txt',f'{total} current executable assertions; fake JDBC and issued auth fixtures.'],
        ['expanded-controls-integration.json',f'{runtime} real app and {sql_checks} SQL/chain checks; isolated synthetic fixtures.'],
        ['expanded-migration.json','Actual additive migration, account identity and grants.'],
        ['legacy-ledger-inventory.json','External local ledger presence/counts; source preservation.'],
        ['dependency-advisory-baseline/current.json','Advisory found in old connector and current live-query result.'],
        ['gitleaks-current.json / repository-controls.json','Current-source scan and custom-rule regressions.'],
        ['integrity-gates.json / archive-scheduler.json','Optimized-mode rejection and actual task/export outcomes.'],
        ['expanded-independent-review.json','Confirmed findings, corrections and final review verdict.'],
        ['ui/security-controls','Six sanitized captures and provenance; injected UI evidence.'],
        ['ui/audit-trail','Read-only business audit dialog states and paging provenance.'],
        ['outputs/all-controls/*.xlsx','Revised working audit worksheet, original source retained.']
    ],[3.05,3.82])
    paragraph(document,'The revised audit workbook preserves the supplied sheets, merged ranges, validations and images while updating assessment fields. It records supported-client/local effectiveness and leaves independent off-host custody Partial. It is a new working copy, not a replacement for the original Milestone 1 evidence.')
    paragraph(document,'References: original Group 13 SIP pages 4 to 7 and audit rows 73 to 83; OWASP Authentication, Authorization and Logging Cheat Sheets; Microsoft DPAPI documentation; MySQL Connector/J compatibility/release documentation; OSV API documentation. Repository security/EXPANDED_CONTROL_PLAN.md and security/README.md link the official sources.')
    document.save(path)
    width,height=landscape(A2)
    c=canvas.Canvas(str(OUTPUT/'MotorPH_Expanded_Control_Diagram.pdf'),pagesize=(width,height))
    c.setTitle('MotorPH Expanded Security Controls')
    c.setFont('Helvetica-Bold',23);c.drawString(45,height-48,'MotorPH expanded security controls')
    c.setFont('Helvetica',12);c.drawString(45,height-72,'Group 13 - isolated local implementation - 17 October 2026')
    image_width=width-90;image_height=image_width*720/1365
    c.drawImage(ImageReader(str(OUTPUT/'MotorPH_Expanded_Control_Diagram.png')),45,height-95-image_height,width=image_width,height=image_height)
    c.setFont('Helvetica',12);y=height-120-image_height
    for text in ['C1 protected credentials/grants/TLS; C2 BCrypt/reset; C3 persistent throttling and HR notification.',
                 'C4 supported-client authorization; C5 transactional audit chain and database workflow ledgers.',
                 'Original database remains unchanged. Local archives are verified; independent off-host custody remains unproven.']:
        c.drawString(45,y,text);y-=19
    c.showPage();c.save()
    print(f'Expanded implementation log authored with {total} offline and {runtime}+{sql_checks} isolated checks.')


if __name__=='__main__':
    main()
